from io import BytesIO

import pandas as pd
from fastapi.testclient import TestClient

from log_anomaly.api import app as api_module
from log_anomaly.api.app import app
from log_anomaly.persistence.sqlite_store import SQLiteStore

client = TestClient(app)


def test_api_health_route() -> None:
    response = client.get("/api/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_health_endpoint() -> None:
    response = client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_analyze_endpoint_uses_shared_pipeline() -> None:
    response = client.post(
        "/api/v1/analyze",
        json={
            "method": "baseline",
            "events": [{"message": "routine heartbeat"}] * 8
            + [{"message": "unexpected database failure"}],
        },
    )

    assert response.status_code == 200
    payload = response.json()["events"]
    assert len(payload) == 9
    assert payload[-1]["is_anomaly"] is True
    assert payload[-1]["explanation"]


def test_analyze_endpoint_rejects_empty_requests() -> None:
    response = client.post("/api/v1/analyze", json={"events": []})

    assert response.status_code == 422


def test_upload_endpoint_parses_analyzes_and_persists_text_file(tmp_path) -> None:
    api_module.store = SQLiteStore(tmp_path / "api.db")
    response = client.post(
        "/api/upload",
        files={
            "file": (
                "sample_logs.txt",
                b"2026-10-28 10:00:01 INFO Application started\n"
                b"2026-10-28 10:00:02 ERROR Database connection timeout\n"
                b"2026-10-28 10:00:03 INFO Cache refreshed\n"
                b"2026-10-28 10:00:04 ERROR Disk corruption detected\n",
                "text/plain",
            )
        },
        data={"method": "baseline"},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["filename"] == "sample_logs.txt"
    assert payload["normalized_encoding"] == "utf-8"
    assert payload["events"]
    assert payload["events"][0]["message"] == "Application started"
    assert "analysis_id" in payload
    assert api_module.store.list_analyses()[0]["filename"] == "sample_logs.txt"


def test_upload_endpoint_rejects_unsupported_pdf_payload(tmp_path) -> None:
    api_module.store = SQLiteStore(tmp_path / "api.db")
    response = client.post(
        "/api/upload",
        files={"file": ("paper.pdf", b"not a PDF", "application/pdf")},
    )

    assert response.status_code == 422
    assert "not a valid PDF" in response.json()["detail"]


def test_history_and_alert_exports_read_persisted_analysis(tmp_path) -> None:
    api_module.store = SQLiteStore(tmp_path / "history.db")
    events = pd.DataFrame(
        {
            "event_id": ["event-1"],
            "group_id": ["INC-0001"],
            "timestamp": [None],
            "severity": ["ERROR"],
            "host": [None],
            "service": [None],
            "message": ["database timeout"],
            "template": ["database timeout"],
            "anomaly_score": [0.9],
            "is_anomaly": [True],
            "explanation": ["test evidence"],
        }
    )
    analysis_id = api_module.store.save_analysis(
        events,
        filename="sample.log",
        file_format="text",
        source_encoding="utf-8",
        method="baseline",
    )

    history = client.get("/api/analyses")
    detail = client.get(f"/api/analyses/{analysis_id}")
    alerts = client.get("/api/alerts")
    exported = client.get("/api/alerts/export?format=csv")

    assert history.status_code == 200
    assert history.json()["analyses"][0]["analysis_id"] == analysis_id
    assert detail.json()["events"][0]["message"] == "database timeout"
    assert alerts.json()["alerts"][0]["analysis_id"] == analysis_id
    assert "database timeout" in exported.text


def test_analysis_exports_create_pdf_and_filtered_excel(tmp_path) -> None:
    api_module.store = SQLiteStore(tmp_path / "report.db")
    events = pd.DataFrame(
        {
            "event_id": ["normal", "anomaly"],
            "group_id": [None, "INC-0001"],
            "timestamp": [None, None],
            "severity": ["INFO", "ERROR"],
            "host": [None, None],
            "service": [None, None],
            "message": ["service started", "database corruption"],
            "template": ["service started", "database corruption"],
            "anomaly_score": [0.0, 1.0],
            "is_anomaly": [False, True],
            "explanation": ["common event", "rare event"],
        }
    )
    analysis_id = api_module.store.save_analysis(
        events,
        filename="sample.log",
        file_format="text",
        source_encoding="utf-8",
        method="baseline",
    )

    pdf = client.get(f"/api/analyses/{analysis_id}/export?format=pdf")
    excel = client.get(f"/api/analyses/{analysis_id}/export?format=xlsx")

    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    from pypdf import PdfReader

    pdf_text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(pdf.content)).pages)
    assert "service started" in pdf_text
    assert "database corruption" not in pdf_text
    assert excel.status_code == 200
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(excel.content), read_only=True)
    assert workbook["Summary"]["B8"].value == 1
    assert workbook["Log records"].max_row == 2
    assert workbook["Log records"]["C2"].value == "service started"


def _text_pdf() -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    page[NameObject("/Resources")] = resources
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 72 720 Td "
        b"(2026-10-28 10:00:00 ERROR database timeout) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_upload_extracts_text_logs(tmp_path) -> None:
    api_module.store = SQLiteStore(tmp_path / "pdf.db")
    response = client.post(
        "/api/upload",
        files={"file": ("deeplog.pdf", _text_pdf(), "application/pdf")},
        data={"method": "baseline"},
    )

    assert response.status_code == 200
    assert response.json()["file_format"].startswith("pdf/")
    assert response.json()["events"][0]["severity"] == "ERROR"


def test_upload_endpoint_normalizes_cp1252_csv(tmp_path) -> None:
    api_module.store = SQLiteStore(tmp_path / "encoding.db")
    response = client.post(
        "/api/upload",
        files={
            "file": (
                "legacy.csv",
                "message;level\ncafé;INFO\n".encode("cp1252"),
                "text/csv",
            )
        },
        data={"method": "baseline"},
    )

    assert response.status_code == 200
    assert response.json()["source_encoding"] == "cp1252"
    assert response.json()["events"][0]["message"] == "café"


def test_upload_endpoint_rejects_scanned_pdf(tmp_path) -> None:
    from pypdf import PdfWriter

    api_module.store = SQLiteStore(tmp_path / "scanned.db")
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    output = BytesIO()
    writer.write(output)
    response = client.post(
        "/api/upload",
        files={"file": ("scanned.pdf", output.getvalue(), "application/pdf")},
    )

    assert response.status_code == 422
    assert "OCR" in response.json()["detail"]


def test_upload_endpoint_rejects_unsupported_file_type() -> None:
    response = client.post(
        "/api/upload",
        files={"file": ("dataset.xlsx", b"arbitrary", "application/octet-stream")},
    )

    assert response.status_code == 415
