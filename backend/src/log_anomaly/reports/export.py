"""Create PDF reports and Excel workbooks for saved analyses."""

from __future__ import annotations

from io import BytesIO
from typing import Any, Literal
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ExportFormat = Literal["pdf", "xlsx"]
MEDIA_TYPES = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def build_analysis_export(
    analysis: dict[str, Any],
    *,
    format: ExportFormat,
    exclude_anomalies: bool,
) -> tuple[bytes, str, str]:
    """Return generated report bytes, media type, and file extension."""
    events = analysis["events"]
    anomalies = [event for event in events if event.get("is_anomaly")]
    output_events = (
        [event for event in events if not event.get("is_anomaly")]
        if exclude_anomalies
        else events
    )
    if format == "pdf":
        content = _build_pdf(analysis, output_events, anomalies, exclude_anomalies)
    else:
        content = _build_excel(analysis, output_events, anomalies, exclude_anomalies)
    return content, MEDIA_TYPES[format], format


def _build_pdf(
    analysis: dict[str, Any],
    events: list[dict[str, Any]],
    anomalies: list[dict[str, Any]],
    exclude_anomalies: bool,
) -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="Anomaly Detect - Log Analysis Report",
    )
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "ReportBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=7,
        leading=9,
        alignment=TA_LEFT,
        wordWrap="CJK",
    )
    heading = styles["Heading2"]
    rows = [[
        "Time", "Severity", "Message", "Host / service", "Score", "Alert group",
    ]]
    for event in events:
        rows.append([
            _paragraph(event.get("timestamp") or "-", body),
            _paragraph(event.get("severity") or "-", body),
            _paragraph(event.get("message") or "", body),
            _paragraph(event.get("host") or event.get("service") or "-", body),
            f"{float(event.get('anomaly_score') or 0):.3f}",
            _paragraph(event.get("group_id") or "-", body),
        ])

    content: list[Any] = [
        Paragraph("Anomaly Detect — Log Analysis Report", styles["Title"]),
        Spacer(1, 4 * mm),
        Paragraph(f"<b>Source file:</b> {_paragraph_text(analysis.get('filename', ''))}", body),
        Paragraph(f"<b>Format:</b> {_paragraph_text(analysis.get('file_format', ''))}", body),
        Paragraph(f"<b>Created:</b> {_paragraph_text(analysis.get('created_at', ''))}", body),
        Spacer(1, 3 * mm),
        Paragraph(
            f"<b>Total records:</b> {len(analysis['events'])} &nbsp; "
            f"<b>Detected anomalies:</b> {len(anomalies)} &nbsp; "
            f"<b>Records in this report:</b> {len(events)} &nbsp; "
            f"<b>Anomalies excluded:</b> {len(anomalies) if exclude_anomalies else 0}",
            body,
        ),
        Spacer(1, 5 * mm),
        Paragraph(
            "Log records" if not exclude_anomalies else "Log records (anomalies removed)",
            heading,
        ),
    ]
    if len(rows) == 1:
        content.append(Paragraph("No log records to include.", body))
    else:
        table = Table(
            rows,
            colWidths=[35 * mm, 20 * mm, 100 * mm, 38 * mm, 17 * mm, 25 * mm],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#183b43")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5d9")),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#f1f5f6")],
                ),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ])
        )
        content.append(table)
    document.build(content)
    return output.getvalue()


def _build_excel(
    analysis: dict[str, Any],
    events: list[dict[str, Any]],
    anomalies: list[dict[str, Any]],
    exclude_anomalies: bool,
) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Anomaly Detect — Log Analysis"])
    summary.append(["Source file", analysis.get("filename", "")])
    summary.append(["Format", analysis.get("file_format", "")])
    summary.append(["Created", analysis.get("created_at", "")])
    summary.append(["Total records", len(analysis["events"])])
    summary.append(["Detected anomalies", len(anomalies)])
    summary.append(["Records included", len(events)])
    summary.append(["Anomalies excluded", len(anomalies) if exclude_anomalies else 0])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 64
    summary["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="183B43")
    summary.merge_cells("A1:B1")

    sheet = workbook.create_sheet("Log records")
    columns = [
        "timestamp", "severity", "message", "host", "service",
        "anomaly_score", "group_id", "explanation",
    ]
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="183B43")
    for event in events:
        sheet.append([event.get(column) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    column_widths = {
        "A": 22, "B": 14, "C": 80, "D": 24,
        "E": 24, "F": 16, "G": 18, "H": 64,
    }
    for column, width in column_widths.items():
        sheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _paragraph(value: object, style: ParagraphStyle) -> Paragraph:
    return Paragraph(_paragraph_text(value), style)


def _paragraph_text(value: object) -> str:
    text = escape(str(value))
    return text.encode("latin-1", errors="replace").decode("latin-1")
